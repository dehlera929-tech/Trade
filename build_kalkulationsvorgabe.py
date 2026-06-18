# -*- coding: utf-8 -*-
"""
Build a merged, bilingual (DE/EN) Sales -> Calculation specification workbook.

Merges:
  * CGoE Uebergabecheckliste V4  (project header, offer type, GCoE services,
    documents, manufacturer reference matrix)
  * Maik - Vorschlag Lieferantenauswahl (per-component manufacturer options
    incl. LV relevance + subcontractors)

Result = one clean, fast-to-fill template. Manufacturer column offers a
per-row dropdown including the option "Frei waehlbar / Free choice".
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- palette ----
NAVY      = "1F3864"   # group / section headers
BLUE      = "2E5496"   # column headers
LIGHTBLUE = "D9E1F2"   # group band
INPUT     = "FFF2CC"   # cells Sales has to fill -> light yellow
INPUTBORD = "BF9000"
ZEBRA     = "F2F5FB"   # alternating rows
WHITE     = "FFFFFF"
GREY      = "808080"
HEADTXT   = "FFFFFF"

FONT = "Calibri"

def f(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

def fill(hexc):
    return PatternFill("solid", fgColor=hexc)

thin = Side(style="thin", color="BFBFBF")
med  = Side(style="medium", color=BLUE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
INBORDER = Border(left=Side(style="thin", color=INPUTBORD),
                  right=Side(style="thin", color=INPUTBORD),
                  top=Side(style="thin", color=INPUTBORD),
                  bottom=Side(style="thin", color=INPUTBORD))

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left", vertical="center", wrap_text=True)
LEFTT  = Alignment(horizontal="left", vertical="top", wrap_text=True)

wb = openpyxl.Workbook()

# ======================================================================
# helper to draw the company-style title banner
# ======================================================================
def banner(ws, last_col, title_de, title_en, sub=""):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    t = ws.cell(1, 1, title_de + "   |   " + title_en)
    t.font = f(16, True, HEADTXT); t.fill = fill(NAVY); t.alignment = LEFT
    s = ws.cell(2, 1, sub)
    s.font = f(9, False, HEADTXT, italic=True); s.fill = fill(BLUE); s.alignment = LEFT
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 16

# ======================================================================
# SHEET 1 - Projekt & Rahmen / Project & Scope
# ======================================================================
ws = wb.active
ws.title = "1 Projekt & Rahmen"
ws.sheet_view.showGridLines = False
last = 5
widths = {"A": 3, "B": 46, "C": 46, "D": 22, "E": 40}
for c, w in widths.items():
    ws.column_dimensions[c].width = w
banner(ws, last, "Kalkulationsvorgabe – Projekt & Rahmen",
       "Calculation Brief – Project & Scope",
       "Vorgabe vom Vertrieb an die Kalkulation  /  Specification from Sales to Calculation")

yesno = DataValidation(type="list", formula1='"Ja / Yes,Nein / No"', allow_blank=True)
ws.add_data_validation(yesno)

r = 4
def section(title_de, title_en):
    global r
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last)
    c = ws.cell(r, 2, f"{title_de}  /  {title_en}")
    c.font = f(11, True, HEADTXT); c.fill = fill(NAVY); c.alignment = LEFT
    ws.row_dimensions[r].height = 20
    r += 1

def field(de, en, value="", input_field=True, dropdown=False, note=""):
    """label (DE) in B, label (EN) in C, value in D, note in E"""
    global r
    b = ws.cell(r, 2, de); b.font = f(10, True); b.alignment = LEFT; b.border = BORDER
    c = ws.cell(r, 3, en); c.font = f(9, False, GREY, italic=True); c.alignment = LEFT; c.border = BORDER
    d = ws.cell(r, 4, value); d.alignment = LEFT
    if input_field:
        d.fill = fill(INPUT); d.border = INBORDER
    else:
        d.border = BORDER
    if dropdown:
        yesno.add(d)
    e = ws.cell(r, 5, note); e.font = f(9, False, GREY, italic=True); e.alignment = LEFT; e.border = BORDER
    ws.row_dimensions[r].height = 18
    r += 1

# --- project head ---
section("Projektdaten", "Project data")
field("Projekt", "Project")
field("Bauvorhaben", "Construction project")
field("Ort", "Location")
field("Baubeginn", "Start of construction")
field("Fertigstellung", "Completion")
field("Gebäudeart", "Type of building", note="z. B. Büro, Schule, Krankenhaus / e.g. office, school, hospital")
field("Unser Kunde", "Our customer")
field("Pre-Sales Mitarbeiter", "Pre-sales staff")
field("Sales Ansprechpartner", "Sales contact")
field("Telefon", "Phone")
field("E-Mail", "E-mail")
field("Bitte zurück bis", "Please return by")
field("Datum der Vorgabe", "Date of brief")

r += 1
# --- offer type ---
section("Angebotsart", "Offer type")
otype = [
 ("Spätere Pauschalierung", "Later flat-rate calculation", True, "z. B. Budgetangebot"),
 ("Öffentliche Vergabe", "Public procurement", True, "Ausschreibungstexte sind einzuhalten / tender texts binding"),
 ("Jour fixe erforderlich", "Jour fixe required", True, ""),
 ("Technische Alternativen möglich", "Technical alternatives possible", True, ""),
 ("Schätzkosten (teilweise) möglich", "Estimated costs (partially) possible", True, ""),
 ("Ist der niedrigste Preis wichtig?", "Does the lowest price matter?", True,
    "Ja = preissensitiv / No = Lösung flexibel"),
 ("Referenzpreise zu vorherigen Angeboten", "Reference prices to previous offers", True,
    "z. B. Rahmenvertrag, Nachtrag zum Hauptangebot"),
 ("Mit Schaltschrank", "With control cabinet", True, ""),
 ("Mit Verkabelung + Feldgerätemontage + el. Anschlüsse",
    "With wiring + field device mounting + el. connections", True, ""),
 ("Nur Feldgerätemontage + el. Anschlüsse",
    "Only field device mounting + el. connections", True, ""),
 ("Nachunternehmerauswahl", "Subcontractor selection", True, ""),
 ("System-/Lösungsplanung erforderlich", "System / solution planning required", True, ""),
]
for de, en, dd, note in otype:
    field(de, en, dropdown=dd, note=note)

r += 1
# --- GCoE services ---
section("Leistungen durch GCoE", "Services provided by GCoE")
svc = [
 ("Auswahl Feldgeräte und Ventile", "Selection of field devices and valves"),
 ("Auswahl Automationsstationen und I/O-Module", "Selection of automation stations and I/O modules"),
 ("Auswertung Preisspiegel Feldgeräte + Ventile", "Price comparison field devices + valves"),
 ("Auswertung Preisspiegel Installation", "Price comparison installation"),
 ("Auswertung Preisspiegel Schaltschränke", "Price comparison control cabinets"),
 ("THF-Bearbeitung komplett", "THF processing complete"),
 ("Besonderheiten 1", "Special features 1"),
 ("Besonderheiten 2", "Special features 2"),
]
for de, en in svc:
    field(de, en, dropdown=True)

ws.freeze_panes = "A4"

# ======================================================================
# SHEET 2 - Kalkulationsvorgabe / Calculation Spec  (the merged core)
# ======================================================================
ws2 = wb.create_sheet("2 Kalkulationsvorgabe")
ws2.sheet_view.showGridLines = False
cols = ["A", "B", "C", "D", "E", "F", "G", "H"]
cw = {"A": 4, "B": 34, "C": 34, "D": 40, "E": 30, "F": 32, "G": 16, "H": 34}
for c, w in cw.items():
    ws2.column_dimensions[c].width = w
banner(ws2, 8, "Kalkulationsvorgabe – Komponenten & Hersteller",
       "Calculation Brief – Components & Manufacturers",
       "Was wie kalkulieren? Hersteller-Vorgabe oder 'Frei wählbar'.  /  "
       "What to calculate how? Manufacturer spec or 'Free choice'.")

# column header row (row 3)
headers = [
    ("Nr.", "No."),
    ("Komponente / Gewerk", "Component / Trade"),
    ("", "(English)"),
    ("Hersteller-Optionen", "Manufacturer options"),
    ("Vorgabe Hersteller", "Manufacturer spec"),
    ("Kalkulationsart", "Calculation mode"),
    ("LV-relevant", "BoQ relevant"),
    ("Bemerkung", "Notes"),
]
hr = 3
for i, (de, en) in enumerate(headers, start=1):
    cell = ws2.cell(hr, i, (de + "\n" + en) if de else en)
    cell.font = f(10, True, HEADTXT); cell.fill = fill(BLUE)
    cell.alignment = CENTER; cell.border = BORDER
ws2.row_dimensions[hr].height = 30

# ---- calculation mode + LV dropdowns (shared) ----
calc_list = ('"Fester Hersteller / Fixed,Preisspiegel / Price comparison,'
             'Frei wählbar / Free choice,Schätzkosten / Estimate,'
             'Nicht kalkulieren / Not calculated"')
dv_calc = DataValidation(type="list", formula1=calc_list, allow_blank=True)
ws2.add_data_validation(dv_calc)
dv_lv = DataValidation(type="list", formula1='"Ja / Yes,Nein / No,Optional"', allow_blank=True)
ws2.add_data_validation(dv_lv)

FREE = "Frei wählbar / Free choice"

# ---- merged component catalogue (from both source files) ----
groups = [
 ("Systemarchitektur", "System architecture", [
   ("MBE / Management- und Bedienebene", "Management & operating level",
        ["ADX", "ADS", "MXI 64", "FREMD / external"]),
   ("Automationsstationen", "Automation stations",
        ["METASYS", "EASY IO", "LOYTEC", "WAGO"]),
   ("Anlagenregler", "System controller",
        ["M4 CGM", "M4 CGE", "ATC", "LOYTEC", "WAGO"]),
   ("I/O-Module", "I/O modules",
        ["M4-XPM", "Romutec 10", "Romutec 30", "Romutec 51", "EAP"]),
   ("Brandschutzklappen-Module (BSK)", "Fire damper modules (BSK)",
        ["Romutec", "Frakta", "RK-Tec", "ST-Steuerung", "Agnosys"]),
 ]),
 ("Feldgeräte", "Field devices", [
   ("Fühler Luft", "Air sensors", ["Johnson Controls", "Thermokon", "S+S"]),
   ("Fühler Wasser", "Water sensors", ["Johnson Controls", "Thermokon", "S+S"]),
   ("Differenzdruck", "Differential pressure", ["Oppermann"]),
   ("Wächter – Frost", "Monitor – frost", ["Johnson Controls", "S+S (bei stetig / if modulating)"]),
   ("Wächter – Differenzdruck", "Monitor – diff. pressure", ["Johnson Controls", "S+S"]),
   ("Wächter – Temperatur", "Monitor – temperature", ["Alre", "S+S"]),
   ("Wächter – Druck (flüssig)", "Monitor – pressure (liquid)", ["Sauter"]),
   ("Wächter – Leckage", "Monitor – leakage", ["Jola", "S+S"]),
   ("Wächter – Rauchmelder", "Monitor – smoke detector", ["Oppermann"]),
   ("Gaswarnanlagen", "Gas warning systems", ["Oppermann", "MCS", "EVD"]),
   ("Wetterstation", "Weather station", ["Thies", "RB Messtechnik", "HKW Wetterprognose / forecast"]),
   ("EX-Geräte inkl. Barriere", "Ex devices incl. barrier", ["Rotork", "Schischek"]),
   ("Klappenantriebe Wasser", "Damper actuators water", ["Johnson Controls", "Belimo"]),
   ("Klappenantriebe Luft", "Damper actuators air", ["Johnson Controls", "Belimo", "IT-AT"]),
   ("Messgeräte Wasser", "Meters water", ["MWK Messsysteme", "Molline"]),
   ("Messgeräte Elektro", "Meters electrical", ["IME (BACnet)", "Janitza (Modbus)"]),
 ]),
 ("Ventile & Antriebe", "Valves & drives", [
   ("Ventile Flansch", "Valves flanged", ["Johnson Controls", "Belimo"]),
   ("Ventile Gewinde", "Valves threaded", ["Johnson Controls", "Belimo", "Frakta"]),
   ("Frequenzumrichter (FU)", "Frequency converter", ["Danfoss"]),
 ]),
 ("IT / Netzwerk / Bedienung", "IT / network / operation", [
   ("Netzwerktechnik", "Network technology", ["IT-AT", "Conrad"]),
   ("PC / Server", "PC / server", ["Dell", "Conrad"]),
   ("Touch-PC Tür", "Touch PC door", ["ICO"]),
   ("Touchpanel BACnet", "Touch panel BACnet", ["Exor"]),
   ("Handebene", "Manual override level", ["Romutec (Tür / door)", "Metz Connect (Hutschiene / DIN rail)"]),
   ("Gateways", "Gateways", ["MBS", "ADFWeb"]),
 ]),
 ("Nachunternehmer", "Subcontractors", [
   ("Installationen", "Installation works",
        ["MFG Gebäudetechnik", "Rendke GmbH", "Käppler Elektrotechnik GmbH"]),
   ("Schaltschrankbau", "Control cabinet manufacturing",
        ["MFG Gebäudetechnik", "Elektro Lehmann", "tech control", "Elbara", "Alltec"]),
 ]),
]

row = hr + 1
nr = 0
for g_de, g_en, items in groups:
    # group band
    ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    gc = ws2.cell(row, 1, f"{g_de}   /   {g_en}")
    gc.font = f(11, True, HEADTXT); gc.fill = fill(NAVY); gc.alignment = LEFT
    ws2.row_dimensions[row].height = 20
    row += 1
    for de, en, opts in items:
        nr += 1
        zebra = ZEBRA if nr % 2 == 0 else WHITE
        # Nr
        a = ws2.cell(row, 1, nr); a.alignment = CENTER; a.border = BORDER; a.fill = fill(zebra); a.font = f(9)
        # Komponente DE
        b = ws2.cell(row, 2, de); b.alignment = LEFT; b.border = BORDER; b.fill = fill(zebra); b.font = f(10, True)
        # Komponente EN
        c = ws2.cell(row, 3, en); c.alignment = LEFT; c.border = BORDER; c.fill = fill(zebra); c.font = f(9, False, GREY, italic=True)
        # options
        d = ws2.cell(row, 4, ", ".join(opts)); d.alignment = LEFT; d.border = BORDER; d.fill = fill(zebra); d.font = f(9)
        # manufacturer spec (input + per-row dropdown)
        e = ws2.cell(row, 5, ""); e.alignment = LEFT; e.fill = fill(INPUT); e.border = INBORDER; e.font = f(10, True)
        choice = opts + [FREE]
        formula = '"' + ",".join(choice) + '"'
        if len(formula) <= 255:
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws2.add_data_validation(dv); dv.add(e)
        # calc mode (input + dropdown)
        ff = ws2.cell(row, 6, ""); ff.alignment = LEFT; ff.fill = fill(INPUT); ff.border = INBORDER; ff.font = f(10)
        dv_calc.add(ff)
        # LV relevant
        g = ws2.cell(row, 7, ""); g.alignment = CENTER; g.fill = fill(INPUT); g.border = INBORDER; g.font = f(10)
        dv_lv.add(g)
        # notes
        h = ws2.cell(row, 8, ""); h.alignment = LEFT; h.fill = fill(INPUT); h.border = INBORDER; h.font = f(10)
        ws2.row_dimensions[row].height = 26
        row += 1

ws2.freeze_panes = "A4"
ws2.sheet_view.showGridLines = False

# legend below table
row += 1
leg = ws2.cell(row, 2,
   "Legende / Legend:  Gelbe Felder = von Sales auszufüllen / yellow cells = to be filled by Sales.   "
   "Kalkulationsart / Calculation mode: Fester Hersteller = fix vorgegeben | Preisspiegel = Angebotsvergleich | "
   "Frei wählbar = Kalkulation entscheidet | Schätzkosten = Budget | Nicht kalkulieren = entfällt.")
leg.font = f(9, False, GREY, italic=True); leg.alignment = LEFTT
ws2.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
ws2.row_dimensions[row].height = 40

# ======================================================================
# SHEET 3 - Dokumente / Documents
# ======================================================================
ws3 = wb.create_sheet("3 Dokumente")
ws3.sheet_view.showGridLines = False
cw3 = {"A": 4, "B": 50, "C": 50, "D": 22, "E": 26, "F": 34}
for c, w in cw3.items():
    ws3.column_dimensions[c].width = w
banner(ws3, 6, "Mitgelieferte Dokumente & Unterlagen",
       "Supplied documents & files",
       "Übergabe an Kalkulation / Bratislava  –  Handover to calculation / Bratislava")

dhead = [("Nr.", "No."), ("Dokument", "Document"), ("", "(English)"),
         ("Mitgeliefert?", "Supplied?"), ("An Bratislava senden?", "Send to Bratislava?"),
         ("Bemerkung", "Notes")]
for i, (de, en) in enumerate(dhead, start=1):
    cell = ws3.cell(3, i, (de + "\n" + en) if de else en)
    cell.font = f(10, True, HEADTXT); cell.fill = fill(BLUE); cell.alignment = CENTER; cell.border = BORDER
ws3.row_dimensions[3].height = 30

dv3a = DataValidation(type="list", formula1='"Ja / Yes,Nein / No,N/A"', allow_blank=True)
dv3b = DataValidation(type="list", formula1='"Ja / Yes,Nein / No"', allow_blank=True)
ws3.add_data_validation(dv3a); ws3.add_data_validation(dv3b)

docs = [
 ("CGoE Übergabecheckliste", "CGoE handover checklist"),
 ("THF-Datei für das Projekt", "THF file for the project"),
 ("Leistungsverzeichnis als PDF", "Specifications as PDF"),
 ("Excel-Liste aus dem THF inkl. Langtexte", "Excel list from THF incl. long texts"),
 ("Auslegung / Festlegung der DDC", "DDC interpretation / definition"),
 ("Angebote der Nachunternehmer", "Subcontractor offers"),
]
rr = 4
for i, (de, en) in enumerate(docs, start=1):
    z = ZEBRA if i % 2 == 0 else WHITE
    a = ws3.cell(rr, 1, i); a.alignment = CENTER; a.border = BORDER; a.fill = fill(z); a.font = f(9)
    b = ws3.cell(rr, 2, de); b.alignment = LEFT; b.border = BORDER; b.fill = fill(z); b.font = f(10, True)
    c = ws3.cell(rr, 3, en); c.alignment = LEFT; c.border = BORDER; c.fill = fill(z); c.font = f(9, False, GREY, italic=True)
    d = ws3.cell(rr, 4, ""); d.fill = fill(INPUT); d.border = INBORDER; dv3a.add(d)
    e = ws3.cell(rr, 5, ""); e.fill = fill(INPUT); e.border = INBORDER; dv3b.add(e)
    g = ws3.cell(rr, 6, ""); g.fill = fill(INPUT); g.border = INBORDER
    ws3.row_dimensions[rr].height = 20
    rr += 1

ws3.freeze_panes = "A4"

# ======================================================================
wb.save("/home/user/Trade/Kalkulationsvorgabe_Sales-Kalkulation_DE-EN.xlsx")
print("saved OK")
